# -*- coding: utf-8 -*-
"""
txt → 番茄「上传章节」可用的 docx

番茄的硬要求（不满足会直接发布失败）：
  · 章节标题必须是 H1 一级标题
  · 标题格式「第N章 标题」，序号用阿拉伯数字
  · 单文件 ≤ 5MB
  · 标题至少 5 个字（连续数字算 1 个字）
  · 全书标题不能重复

用法：
  python build_docx.py              全书
  python build_docx.py 161 175      只做第 161-175 章
"""
import re
import os
import sys
import json
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
CFG_PATH = os.path.join(HERE, "config.json")
if not os.path.exists(CFG_PATH):
    sys.exit("缺少 config.json，请先复制 config.example.json 并填写")

CFG = json.load(open(CFG_PATH, encoding="utf-8"))
SRC = CFG["txt"]
if not os.path.isabs(SRC):
    SRC = os.path.abspath(os.path.join(HERE, "..", SRC))
if not os.path.exists(SRC):
    sys.exit(f"找不到正文文件：{SRC}")

FIX = {int(k): v for k, v in (CFG.get("titleFix") or {}).items()}
TOTAL = int(CFG.get("total") or 0)

lo = int(sys.argv[1]) if len(sys.argv) > 1 else 1
hi = int(sys.argv[2]) if len(sys.argv) > 2 else (TOTAL or 10 ** 9)

raw = open(SRC, encoding="utf-8").read().replace("\r\n", "\n")
CH = re.compile(r"^第(\d+)章[\s\u3000]+(.+)$")
chapters, cur = [], None
VOL = re.compile(r"^第[一二三四五六七八九十百]+卷")   # 卷标题行要丢弃，否则会被当成正文
for ln in raw.split("\n"):
    s = ln.strip()
    if not s:
        continue
    if VOL.match(s):
        continue
    m = CH.match(s)
    if m:
        cur = {"num": int(m.group(1)), "title": m.group(2).strip(), "paras": []}
        chapters.append(cur)
        continue
    if cur:
        cur["paras"].append(s)

print(f"解析到 {len(chapters)} 章")


def width(s):
    """番茄的计数方式：连续数字算 1 个字"""
    return len(re.sub(r"\d+", "X", s))


problems, seen = [], {}
for c in chapters:
    c["final"] = FIX.get(c["num"], c["title"])
    full = f"第{c['num']}章 {c['final']}"
    if width(full) < 5:
        problems.append(f"第{c['num']}章「{c['final']}」只有 {width(full)} 字，番茄要求 ≥5")
    if c["final"] in seen:
        problems.append(f"标题重名：「{c['final']}」出现在第{seen[c['final']]}章和第{c['num']}章")
    seen[c["final"]] = c["num"]

if problems:
    print("\n⚠ 以下问题会导致发布失败，请先在 titleFix 里改名：")
    for p in problems:
        print("  ·", p)
    print()

sel = sorted([c for c in chapters if lo <= c["num"] <= hi], key=lambda c: c["num"])
if not sel:
    sys.exit(f"第 {lo}-{hi} 章没有内容")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "宋体"
st.font.size = Pt(12)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.space_after = Pt(0)

for c in sel:
    h = doc.add_heading(f"第{c['num']}章 {c['final']}", level=1)
    for r in h.runs:
        r.font.name = "宋体"
    for p in c["paras"]:
        doc.add_paragraph(p).alignment = WD_ALIGN_PARAGRAPH.LEFT

name = f"{CFG['bookName']}_第{sel[0]['num']}-{sel[-1]['num']}章.docx"
out = os.path.join(HERE, "..", name)
doc.save(out)
chars = sum(len("".join(c["paras"])) for c in sel)
size = os.path.getsize(out) / 1024 / 1024
print(f"已生成 {name}　{len(sel)} 章　{chars:,} 字符　{size:.2f} MB")
if size > 5:
    print("⚠ 超过 5MB，请拆分上传")
